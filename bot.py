import logging
import os
import re
import shutil
import sqlite3
import tempfile
from datetime import datetime, time
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from telegram import ReplyKeyboardMarkup, Update
from telegram.ext import (
    Application,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

from database import get_connection, init_database

# ============================================================
# الإعدادات (كلها من متغيرات البيئة، ولا شيء ثابت في الكود)
# ============================================================

# على Render تأتي المتغيرات من لوحة التحكم. override=False حتى لا يطغى الملف عليها.
load_dotenv("config.env", override=False)

logging.basicConfig(
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger("shop_bot")

TOKEN = os.getenv("TELEGRAM_TOKEN")
SHOP_NAME = os.getenv("SHOP_NAME", "المتجر")
CURRENCY = os.getenv("CURRENCY", "جنيه")
TZ = ZoneInfo(os.getenv("TIMEZONE", "Africa/Khartoum"))
LOW_STOCK_LIMIT = int(os.getenv("LOW_STOCK_LIMIT", "5"))
BACKUP_HOUR = int(os.getenv("BACKUP_HOUR", "23"))


def parse_ids(raw):
    ids = []
    for part in re.split(r"[,\s]+", raw or ""):
        part = part.strip()
        if part.isdigit():
            ids.append(int(part))
    return ids


# يمكن وضع أكثر من رقم مفصولة بفواصل: المالك أولًا ثم الموظفون
ALLOWED_IDS = parse_ids(os.getenv("OWNER_ID", ""))
ALLOWED = filters.User(user_id=ALLOWED_IDS) if ALLOWED_IDS else None

# ============================================================
# نصوص الأزرار
# ============================================================

BTN_DEBT_NEW = "تسجيل دين جديد"
BTN_DEBT_PAY = "تسديد دين"
BTN_CUSTOMERS = "العملاء"
BTN_DEBTS = "الديون"
BTN_TODAY = "تقرير اليوم"
BTN_MONTH = "تقرير الشهر"
BTN_DELETE = "حذف عميل"
BTN_BACKUP = "نسخ احتياطي"
BTN_CANCEL = "إلغاء"

MENU_TEXTS = [
    BTN_DEBT_NEW, BTN_DEBT_PAY, BTN_CUSTOMERS, BTN_DEBTS,
    BTN_TODAY, BTN_MONTH, BTN_DELETE, BTN_BACKUP, BTN_CANCEL,
]

PAY_CASH = "كاش"
PAY_TRANSFER = "تحويل"
PAY_DEBT = "دين"
PAY_PARTIAL = "دفع جزئي"
MORE_ITEM = "إضافة صنف آخر"
FINISH_SALE = "إنهاء البيع"
WALK_IN = "عميل نقدي"

MENU_FILTER = filters.Regex(
    "^(" + "|".join(re.escape(t) for t in MENU_TEXTS) + ")$"
)
TEXT_STATE = filters.TEXT & ~filters.COMMAND & ~MENU_FILTER

MAIN_KEYBOARD = ReplyKeyboardMarkup(
    [
        [BTN_DEBT_NEW, BTN_DEBT_PAY],
        [BTN_CUSTOMERS, BTN_DEBTS],
        [BTN_TODAY, BTN_MONTH],
        [BTN_DELETE, BTN_BACKUP],
        [BTN_CANCEL],
    ],
    resize_keyboard=True,
)

CANCEL_KEYBOARD = ReplyKeyboardMarkup([[BTN_CANCEL]], resize_keyboard=True)

PAYMENT_KEYBOARD = ReplyKeyboardMarkup(
    [[PAY_CASH, PAY_TRANSFER], [PAY_DEBT, PAY_PARTIAL], [BTN_CANCEL]],
    resize_keyboard=True,
)

CASH_ONLY_KEYBOARD = ReplyKeyboardMarkup(
    [[PAY_CASH, PAY_TRANSFER], [BTN_CANCEL]],
    resize_keyboard=True,
)

# حالات المحادثات المستخدمة في نسخة إدارة الديون
DEBT_NAME, DEBT_AMOUNT, DEBT_METHOD, DELETE_NAME, DELETE_CONFIRM = range(5)
NEW_DEBT_NAME, NEW_DEBT_ITEM, NEW_DEBT_AMOUNT = range(5, 8)

# ============================================================
# أدوات عامة
# ============================================================

ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "0123456789" * 2)


def now():
    return datetime.now(TZ).strftime("%Y-%m-%d %H:%M:%S")


def normalize_name(text):
    return " ".join(text.split())


def normalize_number_text(text):
    text = text.translate(ARABIC_DIGITS)
    return text.replace("٬", "").replace(",", "").replace("٫", ".").strip()


def parse_amount(text):
    """عدد صحيح موجب أو صفر، يقبل الأرقام العربية والفواصل."""
    value = normalize_number_text(text)
    if re.fullmatch(r"\d+", value):
        return int(value)
    return None


def parse_quantity(text):
    """كمية قد تكون كسرية (كيلو، لتر)."""
    value = normalize_number_text(text)
    if not re.fullmatch(r"\d+(\.\d+)?", value):
        return None
    number = round(float(value), 3)
    return int(number) if number.is_integer() else number


def format_money(amount):
    return f"{int(round(amount)):,}"


def money(amount):
    return f"{format_money(amount)} {CURRENCY}"


def format_qty(quantity):
    quantity = float(quantity)
    return str(int(quantity)) if quantity.is_integer() else f"{quantity:g}"


async def send_long(update, header, lines, keyboard=None):
    """تيليجرام يرفض الرسائل الأطول من 4096 حرفًا، لذلك نقسمها."""
    chunk = header
    for line in lines:
        if len(chunk) + len(line) + 1 > 3500:
            await update.message.reply_text(chunk)
            chunk = ""
        chunk += ("\n" if chunk else "") + line
    if chunk:
        await update.message.reply_text(chunk, reply_markup=keyboard)


# ============================================================
# العملاء والديون
# ============================================================

def ensure_debt_details_table():
    """يحفظ وصف السلعة مع سجل الدين دون ربطه بجدول المنتجات أو المخزون."""
    connection = get_connection()
    try:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS debt_details (
                sale_id INTEGER PRIMARY KEY,
                product_description TEXT NOT NULL,
                amount NUMERIC NOT NULL,
                created_at TEXT NOT NULL
            )
            """
        )
        connection.commit()
    finally:
        connection.close()


def register_direct_debt(customer_id, customer_name, description, amount):
    """يسجل الدين ووصفه، وينشئ العميل الجديد في المعاملة نفسها عند الحاجة."""
    connection = get_connection()
    try:
        cursor = connection.cursor()
        cursor.execute("BEGIN IMMEDIATE")
        created_at = now()
        if customer_id is None:
            row = cursor.execute(
                "SELECT id FROM customers WHERE name = ?", (customer_name,)
            ).fetchone()
            if row:
                customer_id = row["id"]
            else:
                cursor.execute(
                    "INSERT INTO customers (name, created_at) VALUES (?, ?)",
                    (customer_name, created_at),
                )
                customer_id = cursor.lastrowid
        cursor.execute(
            """
            INSERT INTO sales (
                customer_id, total_amount, paid_amount, debt_amount,
                payment_method, created_at
            ) VALUES (?, ?, 0, ?, ?, ?)
            """,
            (customer_id, amount, amount, PAY_DEBT, created_at),
        )
        sale_id = cursor.lastrowid
        cursor.execute(
            """
            INSERT INTO debt_details (sale_id, product_description, amount, created_at)
            VALUES (?, ?, ?, ?)
            """,
            (sale_id, description, amount, created_at),
        )
        connection.commit()
        return sale_id, customer_id
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()

def get_customer_debt(customer_id):
    connection = get_connection()
    try:
        sales_debt = connection.execute(
            "SELECT COALESCE(SUM(debt_amount), 0) FROM sales WHERE customer_id = ?",
            (customer_id,),
        ).fetchone()[0]
        payments = connection.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM payments WHERE customer_id = ?",
            (customer_id,),
        ).fetchone()[0]
    finally:
        connection.close()
    return max(sales_debt - payments, 0)


DEBT_SQL = """
    SELECT
        c.id,
        c.name,
        COALESCE((SELECT SUM(debt_amount) FROM sales WHERE customer_id = c.id), 0)
        - COALESCE((SELECT SUM(amount) FROM payments WHERE customer_id = c.id), 0)
        AS debt
    FROM customers c
    ORDER BY c.name
"""


def all_customer_debts():
    connection = get_connection()
    try:
        rows = connection.execute(DEBT_SQL).fetchall()
    finally:
        connection.close()
    return [(r["id"], r["name"], max(r["debt"], 0)) for r in rows]


def get_customer_by_name(name):
    connection = get_connection()
    try:
        return connection.execute(
            "SELECT id, name FROM customers WHERE name = ?",
            (normalize_name(name),),
        ).fetchone()
    finally:
        connection.close()


def find_customers(text):
    """بحث بالاسم الكامل أولًا ثم بجزء من الاسم."""
    text = normalize_name(text)
    connection = get_connection()
    try:
        rows = connection.execute(
            "SELECT id, name FROM customers WHERE name = ?", (text,)
        ).fetchall()
        if rows:
            return rows
        return connection.execute(
            "SELECT id, name FROM customers WHERE name LIKE ? ORDER BY name LIMIT 10",
            (f"%{text}%",),
        ).fetchall()
    finally:
        connection.close()


def add_payment(customer_id, amount, method):
    connection = get_connection()
    try:
        connection.execute(
            """
            INSERT INTO payments (customer_id, amount, payment_method, created_at)
            VALUES (?, ?, ?, ?)
            """,
            (customer_id, amount, method, now()),
        )
        connection.commit()
    finally:
        connection.close()


# ============================================================
# المنتجات والمخزون
# ============================================================

PRODUCT_COLUMNS = "id, name, price, cost, stock"


def find_products(text):
    text = normalize_name(text)
    digits = normalize_number_text(text)
    connection = get_connection()
    try:
        if digits.isdigit():
            rows = connection.execute(
                f"SELECT {PRODUCT_COLUMNS} FROM products WHERE id = ? AND is_active = 1",
                (int(digits),),
            ).fetchall()
            if rows:
                return rows

        rows = connection.execute(
            f"SELECT {PRODUCT_COLUMNS} FROM products WHERE is_active = 1 AND name = ?",
            (text,),
        ).fetchall()
        if rows:
            return rows

        return connection.execute(
            f"""
            SELECT {PRODUCT_COLUMNS} FROM products
            WHERE is_active = 1 AND name LIKE ?
            ORDER BY name LIMIT 10
            """,
            (f"%{text}%",),
        ).fetchall()
    finally:
        connection.close()


def product_line(row):
    return (
        f"{row['id']}. {row['name']} — {money(row['price'])} "
        f"— المتوفر: {format_qty(row['stock'])}"
    )


async def choose_product(update, text, pick_state, retry_state):
    """يعيد (منتج، None) عند النجاح، أو (None، الحالة التالية) عند الفشل أو التعدد."""
    matches = find_products(text)

    if not matches:
        await update.message.reply_text(
            "لم أجد منتجًا بهذا الاسم. اكتب جزءًا من الاسم أو رقم المنتج."
        )
        return None, retry_state

    if len(matches) > 1:
        await update.message.reply_text(
            "وجدت أكثر من منتج، اكتب رقم المنتج المطلوب:\n\n"
            + "\n".join(product_line(r) for r in matches)
        )
        return None, pick_state

    return matches[0], None


# ============================================================
# تسجيل البيع (سلة متعددة الأصناف)
# ============================================================

def cart_total(cart):
    return sum(item["line_total"] for item in cart)


def cart_text(cart):
    lines = [
        f"{item['name']} × {format_qty(item['quantity'])} = {format_money(item['line_total'])}"
        for item in cart
    ]
    return "\n".join(lines) + f"\n\nالإجمالي: {money(cart_total(cart))}"


def finalize_sale(customer_name, cart, method, paid):
    """يسجل الفاتورة وينقص المخزون في معاملة واحدة. أي خطأ يلغي كل شيء."""
    connection = get_connection()
    try:
        cursor = connection.cursor()
        cursor.execute("BEGIN IMMEDIATE")

        total = 0
        prepared = []

        for item in cart:
            row = cursor.execute(
                "SELECT name, cost, stock FROM products WHERE id = ? AND is_active = 1",
                (item["product_id"],),
            ).fetchone()

            if not row:
                raise ValueError(f"المنتج {item['name']} لم يعد متاحًا.")

            if row["stock"] + 1e-9 < item["quantity"]:
                raise ValueError(
                    f"المخزون غير كافٍ للمنتج {row['name']} "
                    f"(المتوفر {format_qty(row['stock'])})."
                )

            line_total = int(round(item["quantity"] * item["unit_price"]))
            total += line_total
            prepared.append((item, row["cost"], line_total))

        if method in (PAY_CASH, PAY_TRANSFER):
            paid = total
        elif method == PAY_DEBT:
            paid = 0
        elif not (0 < paid < total):
            raise ValueError("المبلغ المدفوع جزئيًا يجب أن يكون أقل من الإجمالي.")

        debt = total - paid

        customer_id = None
        if customer_name:
            row = cursor.execute(
                "SELECT id FROM customers WHERE name = ?", (customer_name,)
            ).fetchone()
            if row:
                customer_id = row["id"]
            else:
                cursor.execute(
                    "INSERT INTO customers (name, created_at) VALUES (?, ?)",
                    (customer_name, now()),
                )
                customer_id = cursor.lastrowid

        if debt > 0 and customer_id is None:
            raise ValueError("الدين يحتاج إلى اسم عميل.")

        cursor.execute(
            """
            INSERT INTO sales (
                customer_id, total_amount, paid_amount, debt_amount,
                payment_method, created_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (customer_id, total, paid, debt, method, now()),
        )
        sale_id = cursor.lastrowid

        low_stock = []

        for item, cost, line_total in prepared:
            cursor.execute(
                """
                INSERT INTO sale_items (
                    sale_id, product_id, product_name, quantity,
                    unit_price, unit_cost, line_total
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    sale_id, item["product_id"], item["name"], item["quantity"],
                    item["unit_price"], cost, line_total,
                ),
            )
            cursor.execute(
                "UPDATE products SET stock = stock - ? WHERE id = ?",
                (item["quantity"], item["product_id"]),
            )
            cursor.execute(
                """
                INSERT INTO stock_movements (
                    product_id, quantity_change, reason, created_at
                ) VALUES (?, ?, ?, ?)
                """,
                (item["product_id"], -item["quantity"], f"بيع رقم {sale_id}", now()),
            )
            left = cursor.execute(
                "SELECT stock FROM products WHERE id = ?", (item["product_id"],)
            ).fetchone()["stock"]
            if left <= LOW_STOCK_LIMIT:
                low_stock.append((item["name"], left))

        connection.commit()

        return {
            "sale_id": sale_id,
            "total": total,
            "paid": paid,
            "debt": debt,
            "customer_id": customer_id,
            "low_stock": low_stock,
        }

    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


async def sale_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()

    connection = get_connection()
    try:
        count = connection.execute(
            "SELECT COUNT(*) FROM products WHERE is_active = 1"
        ).fetchone()[0]
    finally:
        connection.close()

    if count == 0:
        await update.message.reply_text(
            "لا توجد منتجات مسجلة. أضف منتجًا أولًا من زر إضافة منتج.",
            reply_markup=MAIN_KEYBOARD,
        )
        return ConversationHandler.END

    context.user_data["cart"] = []

    await update.message.reply_text(
        "اكتب اسم المنتج (أو جزءًا منه) أو رقمه:",
        reply_markup=CANCEL_KEYBOARD,
    )
    return SALE_ITEM


async def sale_item(update: Update, context: ContextTypes.DEFAULT_TYPE):
    product, retry = await choose_product(
        update, update.message.text, SALE_PICK, SALE_ITEM
    )
    if product is None:
        return retry

    if product["stock"] <= 0:
        await update.message.reply_text(
            f"المنتج {product['name']} نفد من المخزون. اختر منتجًا آخر."
        )
        return SALE_ITEM

    context.user_data["pending"] = {
        "product_id": product["id"],
        "name": product["name"],
        "price": product["price"],
        "stock": product["stock"],
    }

    await update.message.reply_text(
        f"{product['name']}\n"
        f"السعر: {money(product['price'])}\n"
        f"المتوفر: {format_qty(product['stock'])}\n\n"
        "أدخل الكمية:"
    )
    return SALE_QTY


async def sale_qty(update: Update, context: ContextTypes.DEFAULT_TYPE):
    quantity = parse_quantity(update.message.text)

    if quantity is None or quantity <= 0:
        await update.message.reply_text("أدخل كمية صحيحة أكبر من صفر.")
        return SALE_QTY

    pending = context.user_data["pending"]
    cart = context.user_data["cart"]

    in_cart = sum(i["quantity"] for i in cart if i["product_id"] == pending["product_id"])

    if quantity + in_cart > pending["stock"] + 1e-9:
        await update.message.reply_text(
            f"المتوفر {format_qty(pending['stock'])} فقط "
            f"(في السلة منه {format_qty(in_cart)}). أدخل كمية أقل."
        )
        return SALE_QTY

    existing = next(
        (i for i in cart if i["product_id"] == pending["product_id"]), None
    )

    if existing:
        existing["quantity"] = round(existing["quantity"] + quantity, 3)
        existing["line_total"] = int(round(existing["quantity"] * existing["unit_price"]))
    else:
        cart.append(
            {
                "product_id": pending["product_id"],
                "name": pending["name"],
                "quantity": quantity,
                "unit_price": pending["price"],
                "line_total": int(round(quantity * pending["price"])),
            }
        )

    context.user_data.pop("pending", None)

    await update.message.reply_text(
        "السلة الحالية:\n\n" + cart_text(cart),
        reply_markup=MORE_KEYBOARD,
    )
    return SALE_MORE


async def sale_more(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()

    if text == MORE_ITEM:
        await update.message.reply_text(
            "اكتب اسم المنتج أو رقمه:", reply_markup=CANCEL_KEYBOARD
        )
        return SALE_ITEM

    if text == FINISH_SALE:
        await update.message.reply_text(
            "اكتب اسم العميل، أو اضغط عميل نقدي إن لم يكن له حساب:",
            reply_markup=CUSTOMER_KEYBOARD,
        )
        return SALE_CUSTOMER

    await update.message.reply_text("اختر من الأزرار.")
    return SALE_MORE


async def sale_customer(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    total = cart_total(context.user_data["cart"])

    if text == WALK_IN:
        context.user_data["customer_name"] = None
        await update.message.reply_text(
            f"الإجمالي: {money(total)}\n\nاختر طريقة الدفع:",
            reply_markup=CASH_ONLY_KEYBOARD,
        )
        return SALE_PAYMENT

    name = normalize_name(text)
    context.user_data["customer_name"] = name

    customer = get_customer_by_name(name)

    if customer:
        debt = get_customer_debt(customer["id"])
        note = f"عميل مسجل. دينه الحالي: {money(debt)}"
    else:
        note = "عميل جديد وسيتم تسجيله."

    await update.message.reply_text(
        f"العميل: {name}\n{note}\n\n"
        f"الإجمالي: {money(total)}\n\nاختر طريقة الدفع:",
        reply_markup=PAYMENT_KEYBOARD,
    )
    return SALE_PAYMENT


async def sale_payment(update: Update, context: ContextTypes.DEFAULT_TYPE):
    method = update.message.text.strip()
    walk_in = context.user_data.get("customer_name") is None
    allowed = [PAY_CASH, PAY_TRANSFER] if walk_in else [
        PAY_CASH, PAY_TRANSFER, PAY_DEBT, PAY_PARTIAL
    ]

    if method not in allowed:
        await update.message.reply_text("اختر إحدى طرق الدفع الموجودة في القائمة.")
        return SALE_PAYMENT

    if method == PAY_PARTIAL:
        total = cart_total(context.user_data["cart"])
        context.user_data["method"] = method
        await update.message.reply_text(
            f"الإجمالي: {money(total)}\n\nأدخل المبلغ المدفوع الآن:"
        )
        return SALE_PAID

    return await complete_sale(update, context, method, None)


async def sale_paid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    paid = parse_amount(update.message.text)
    total = cart_total(context.user_data["cart"])

    if paid is None or paid <= 0:
        await update.message.reply_text("أدخل مبلغًا صحيحًا أكبر من صفر.")
        return SALE_PAID

    if paid >= total:
        await update.message.reply_text(
            f"المبلغ يجب أن يكون أقل من الإجمالي ({money(total)}). "
            "للدفع الكامل اضغط إلغاء ثم اختر كاش أو تحويل."
        )
        return SALE_PAID

    return await complete_sale(update, context, PAY_PARTIAL, paid)


async def complete_sale(update, context, method, paid):
    cart = context.user_data["cart"]
    customer_name = context.user_data.get("customer_name")

    try:
        result = finalize_sale(customer_name, cart, method, paid)
    except ValueError as error:
        await update.message.reply_text(
            f"تعذر تسجيل البيع: {error}", reply_markup=MAIN_KEYBOARD
        )
        context.user_data.clear()
        return ConversationHandler.END
    except Exception:
        logger.exception("sale failed")
        await update.message.reply_text(
            "حدث خطأ غير متوقع ولم يتم تسجيل البيع. حاول مرة أخرى.",
            reply_markup=MAIN_KEYBOARD,
        )
        context.user_data.clear()
        return ConversationHandler.END

    lines = [
        f"تم تسجيل الفاتورة رقم {result['sale_id']}",
        "",
        cart_text(cart),
        "",
        f"المدفوع: {money(result['paid'])}",
        f"المتبقي: {money(result['debt'])}",
        f"طريقة الدفع: {method}",
    ]

    if customer_name:
        lines.append(f"العميل: {customer_name}")
        lines.append(
            f"إجمالي دينه الآن: {money(get_customer_debt(result['customer_id']))}"
        )

    if result["low_stock"]:
        lines.append("")
        for name, left in result["low_stock"]:
            lines.append(f"تنبيه: {name} قارب على النفاد، المتبقي {format_qty(left)}")

    await update.message.reply_text("\n".join(lines), reply_markup=MAIN_KEYBOARD)
    context.user_data.clear()
    return ConversationHandler.END


# ============================================================
# تسجيل دين جديد
# ============================================================

async def new_debt_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text(
        "اكتب اسم العميل كاملًا:", reply_markup=CANCEL_KEYBOARD
    )
    return NEW_DEBT_NAME


async def new_debt_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    name = normalize_name(update.message.text)
    if not name:
        await update.message.reply_text("اكتب اسم العميل.")
        return NEW_DEBT_NAME

    customer = get_customer_by_name(name)
    if customer:
        customer_id = customer["id"]
        saved_name = customer["name"]
    else:
        customer_id = None
        saved_name = name

    context.user_data["new_debt_customer_id"] = customer_id
    context.user_data["new_debt_customer_name"] = saved_name
    await update.message.reply_text(
        f"العميل: {saved_name}\nاكتب اسم المنتج أو وصف ما أخذه:"
    )
    return NEW_DEBT_ITEM


async def new_debt_item(update: Update, context: ContextTypes.DEFAULT_TYPE):
    description = normalize_name(update.message.text)
    if not description:
        await update.message.reply_text("اكتب اسم المنتج أو وصفه.")
        return NEW_DEBT_ITEM

    context.user_data["new_debt_description"] = description
    await update.message.reply_text("اكتب قيمة الدين:")
    return NEW_DEBT_AMOUNT


async def new_debt_amount(update: Update, context: ContextTypes.DEFAULT_TYPE):
    amount = parse_amount(update.message.text)
    if amount is None or amount <= 0:
        await update.message.reply_text("أدخل مبلغًا صحيحًا أكبر من صفر.")
        return NEW_DEBT_AMOUNT

    data = context.user_data
    try:
        sale_id, customer_id = register_direct_debt(
            data["new_debt_customer_id"],
            data["new_debt_customer_name"],
            data["new_debt_description"],
            amount,
        )
        balance = get_customer_debt(customer_id)
    except Exception as error:
        logger.exception("register direct debt failed")
        await update.message.reply_text(
            f"تعذر تسجيل الدين: {error}", reply_markup=MAIN_KEYBOARD
        )
        context.user_data.clear()
        return ConversationHandler.END

    await update.message.reply_text(
        "تم تسجيل الدين بنجاح.\n\n"
        f"رقم السجل: {sale_id}\n"
        f"العميل: {data['new_debt_customer_name']}\n"
        f"المنتج/الوصف: {data['new_debt_description']}\n"
        f"المبلغ: {money(amount)}\n"
        f"التاريخ: {now()}\n"
        f"إجمالي دين العميل الآن: {money(balance)}",
        reply_markup=MAIN_KEYBOARD,
    )
    context.user_data.clear()
    return ConversationHandler.END


# ============================================================
# تسديد دين
# ============================================================

async def debt_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text(
        "اكتب اسم العميل الذي يسدد الدين:", reply_markup=CANCEL_KEYBOARD
    )
    return DEBT_NAME


async def debt_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    matches = find_customers(update.message.text)

    if not matches:
        await update.message.reply_text("هذا العميل غير موجود. جرّب جزءًا آخر من الاسم.")
        return DEBT_NAME

    if len(matches) > 1:
        await update.message.reply_text(
            "وجدت أكثر من عميل، اكتب الاسم كاملًا:\n\n"
            + "\n".join(r["name"] for r in matches)
        )
        return DEBT_NAME

    customer = matches[0]
    debt = get_customer_debt(customer["id"])

    if debt <= 0:
        await update.message.reply_text(
            f"العميل {customer['name']} ليس عليه دين حاليًا.",
            reply_markup=MAIN_KEYBOARD,
        )
        context.user_data.clear()
        return ConversationHandler.END

    context.user_data["debt_customer_id"] = customer["id"]
    context.user_data["debt_customer_name"] = customer["name"]
    context.user_data["debt_balance"] = debt

    await update.message.reply_text(
        f"العميل: {customer['name']}\n"
        f"الدين الحالي: {money(debt)}\n\n"
        "أدخل المبلغ المسدد:"
    )
    return DEBT_AMOUNT


async def debt_amount(update: Update, context: ContextTypes.DEFAULT_TYPE):
    amount = parse_amount(update.message.text)
    debt = context.user_data["debt_balance"]

    if amount is None or amount <= 0:
        await update.message.reply_text("أدخل مبلغًا صحيحًا أكبر من صفر.")
        return DEBT_AMOUNT

    if amount > debt:
        await update.message.reply_text(
            f"لا يمكن تسديد أكثر من الدين الحالي ({money(debt)})."
        )
        return DEBT_AMOUNT

    context.user_data["debt_payment_amount"] = amount

    await update.message.reply_text(
        "اختر طريقة التسديد:", reply_markup=CASH_ONLY_KEYBOARD
    )
    return DEBT_METHOD


async def debt_method(update: Update, context: ContextTypes.DEFAULT_TYPE):
    method = update.message.text.strip()

    if method not in (PAY_CASH, PAY_TRANSFER):
        await update.message.reply_text("اختر كاش أو تحويل.")
        return DEBT_METHOD

    data = context.user_data
    add_payment(data["debt_customer_id"], data["debt_payment_amount"], method)
    new_debt = get_customer_debt(data["debt_customer_id"])

    await update.message.reply_text(
        "تم تسجيل التسديد.\n\n"
        f"العميل: {data['debt_customer_name']}\n"
        f"المبلغ المسدد: {money(data['debt_payment_amount'])}\n"
        f"طريقة التسديد: {method}\n"
        f"الدين قبل التسديد: {money(data['debt_balance'])}\n"
        f"الدين المتبقي: {money(new_debt)}",
        reply_markup=MAIN_KEYBOARD,
    )

    context.user_data.clear()
    return ConversationHandler.END


# ============================================================
# إضافة منتج
# ============================================================

async def product_add_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text("اكتب اسم المنتج الجديد:", reply_markup=CANCEL_KEYBOARD)
    return PROD_NAME


async def product_add_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    name = normalize_name(update.message.text)

    connection = get_connection()
    try:
        exists = connection.execute(
            "SELECT 1 FROM products WHERE name = ?", (name,)
        ).fetchone()
    finally:
        connection.close()

    if exists:
        await update.message.reply_text(
            "هذا المنتج مسجل مسبقًا. استخدم توريد مخزون أو تعديل سعر.",
            reply_markup=MAIN_KEYBOARD,
        )
        context.user_data.clear()
        return ConversationHandler.END

    context.user_data["new_product_name"] = name
    await update.message.reply_text(f"سعر بيع الوحدة بالـ{CURRENCY}:")
    return PROD_PRICE


async def product_add_price(update: Update, context: ContextTypes.DEFAULT_TYPE):
    price = parse_amount(update.message.text)

    if price is None or price <= 0:
        await update.message.reply_text("أدخل سعرًا صحيحًا أكبر من صفر.")
        return PROD_PRICE

    context.user_data["new_product_price"] = price
    await update.message.reply_text(
        "سعر شراء الوحدة (التكلفة) لحساب الربح. اكتب 0 إن لم تعرفه:"
    )
    return PROD_COST


async def product_add_cost(update: Update, context: ContextTypes.DEFAULT_TYPE):
    cost = parse_amount(update.message.text)

    if cost is None:
        await update.message.reply_text("أدخل رقمًا صحيحًا.")
        return PROD_COST

    context.user_data["new_product_cost"] = cost
    await update.message.reply_text("الكمية الموجودة الآن في المخزون:")
    return PROD_STOCK


async def product_add_stock(update: Update, context: ContextTypes.DEFAULT_TYPE):
    stock = parse_quantity(update.message.text)

    if stock is None:
        await update.message.reply_text("أدخل كمية صحيحة (صفر أو أكثر).")
        return PROD_STOCK

    data = context.user_data
    connection = get_connection()
    try:
        cursor = connection.cursor()
        cursor.execute(
            """
            INSERT INTO products (name, price, cost, stock, created_at)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                data["new_product_name"], data["new_product_price"],
                data["new_product_cost"], stock, now(),
            ),
        )
        product_id = cursor.lastrowid
        if stock > 0:
            cursor.execute(
                """
                INSERT INTO stock_movements (
                    product_id, quantity_change, reason, created_at
                ) VALUES (?, ?, ?, ?)
                """,
                (product_id, stock, "رصيد افتتاحي", now()),
            )
        connection.commit()
    finally:
        connection.close()

    await update.message.reply_text(
        "تمت إضافة المنتج.\n\n"
        f"{product_id}. {data['new_product_name']}\n"
        f"سعر البيع: {money(data['new_product_price'])}\n"
        f"التكلفة: {money(data['new_product_cost'])}\n"
        f"المخزون: {format_qty(stock)}",
        reply_markup=MAIN_KEYBOARD,
    )

    context.user_data.clear()
    return ConversationHandler.END


# ============================================================
# توريد مخزون
# ============================================================

async def restock_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text(
        "اكتب اسم المنتج الموردة كميته أو رقمه:", reply_markup=CANCEL_KEYBOARD
    )
    return RESTOCK_ITEM


async def restock_item(update: Update, context: ContextTypes.DEFAULT_TYPE):
    product, retry = await choose_product(
        update, update.message.text, RESTOCK_PICK, RESTOCK_ITEM
    )
    if product is None:
        return retry

    context.user_data["restock_id"] = product["id"]
    context.user_data["restock_name"] = product["name"]
    context.user_data["restock_old_cost"] = product["cost"]

    await update.message.reply_text(
        f"{product['name']}\n"
        f"المخزون الحالي: {format_qty(product['stock'])}\n\n"
        "أدخل الكمية الموردة:"
    )
    return RESTOCK_QTY


async def restock_qty(update: Update, context: ContextTypes.DEFAULT_TYPE):
    quantity = parse_quantity(update.message.text)

    if quantity is None or quantity <= 0:
        await update.message.reply_text("أدخل كمية صحيحة أكبر من صفر.")
        return RESTOCK_QTY

    context.user_data["restock_qty"] = quantity
    await update.message.reply_text(
        f"سعر شراء الوحدة الجديد. التكلفة المسجلة الآن "
        f"{money(context.user_data['restock_old_cost'])}.\n"
        "اكتب 0 للإبقاء عليها دون تغيير:"
    )
    return RESTOCK_COST


async def restock_cost(update: Update, context: ContextTypes.DEFAULT_TYPE):
    cost = parse_amount(update.message.text)

    if cost is None:
        await update.message.reply_text("أدخل رقمًا صحيحًا.")
        return RESTOCK_COST

    data = context.user_data
    connection = get_connection()
    try:
        cursor = connection.cursor()
        cursor.execute(
            "UPDATE products SET stock = stock + ? WHERE id = ?",
            (data["restock_qty"], data["restock_id"]),
        )
        if cost > 0:
            cursor.execute(
                "UPDATE products SET cost = ? WHERE id = ?",
                (cost, data["restock_id"]),
            )
        cursor.execute(
            """
            INSERT INTO stock_movements (
                product_id, quantity_change, reason, created_at
            ) VALUES (?, ?, ?, ?)
            """,
            (data["restock_id"], data["restock_qty"], "توريد", now()),
        )
        new_stock = cursor.execute(
            "SELECT stock FROM products WHERE id = ?", (data["restock_id"],)
        ).fetchone()["stock"]
        connection.commit()
    finally:
        connection.close()

    await update.message.reply_text(
        "تم تسجيل التوريد.\n\n"
        f"المنتج: {data['restock_name']}\n"
        f"الكمية المضافة: {format_qty(data['restock_qty'])}\n"
        f"المخزون الآن: {format_qty(new_stock)}",
        reply_markup=MAIN_KEYBOARD,
    )

    context.user_data.clear()
    return ConversationHandler.END


# ============================================================
# تعديل سعر البيع
# ============================================================

async def price_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text(
        "اكتب اسم المنتج المراد تعديل سعره أو رقمه:", reply_markup=CANCEL_KEYBOARD
    )
    return PRICE_ITEM


async def price_item(update: Update, context: ContextTypes.DEFAULT_TYPE):
    product, retry = await choose_product(
        update, update.message.text, PRICE_PICK, PRICE_ITEM
    )
    if product is None:
        return retry

    context.user_data["price_id"] = product["id"]
    context.user_data["price_name"] = product["name"]
    context.user_data["price_old"] = product["price"]

    await update.message.reply_text(
        f"{product['name']}\n"
        f"السعر الحالي: {money(product['price'])}\n\n"
        "أدخل السعر الجديد:"
    )
    return PRICE_VALUE


async def price_value(update: Update, context: ContextTypes.DEFAULT_TYPE):
    price = parse_amount(update.message.text)

    if price is None or price <= 0:
        await update.message.reply_text("أدخل سعرًا صحيحًا أكبر من صفر.")
        return PRICE_VALUE

    data = context.user_data
    connection = get_connection()
    try:
        connection.execute(
            "UPDATE products SET price = ? WHERE id = ?",
            (price, data["price_id"]),
        )
        connection.commit()
    finally:
        connection.close()

    await update.message.reply_text(
        f"تم تعديل سعر {data['price_name']}\n"
        f"من {money(data['price_old'])} إلى {money(price)}\n\n"
        "الفواتير السابقة لا تتغير.",
        reply_markup=MAIN_KEYBOARD,
    )

    context.user_data.clear()
    return ConversationHandler.END


# ============================================================
# القوائم السريعة
# ============================================================

async def products_list(update: Update, context: ContextTypes.DEFAULT_TYPE):
    connection = get_connection()
    try:
        rows = connection.execute(
            f"SELECT {PRODUCT_COLUMNS} FROM products WHERE is_active = 1 ORDER BY name"
        ).fetchall()
    finally:
        connection.close()

    if not rows:
        await update.message.reply_text("لا توجد منتجات مسجلة.")
        return

    await send_long(update, "المنتجات:\n", [product_line(r) for r in rows])


async def low_stock_list(update: Update, context: ContextTypes.DEFAULT_TYPE):
    connection = get_connection()
    try:
        rows = connection.execute(
            f"""
            SELECT {PRODUCT_COLUMNS} FROM products
            WHERE is_active = 1 AND stock <= ?
            ORDER BY stock, name
            """,
            (LOW_STOCK_LIMIT,),
        ).fetchall()
    finally:
        connection.close()

    if not rows:
        await update.message.reply_text("لا توجد منتجات قاربت على النفاد.")
        return

    await send_long(
        update,
        f"منتجات مخزونها {LOW_STOCK_LIMIT} أو أقل:\n",
        [f"{r['name']} — المتبقي: {format_qty(r['stock'])}" for r in rows],
    )


async def customers_list(update: Update, context: ContextTypes.DEFAULT_TYPE):
    rows = all_customer_debts()

    if not rows:
        await update.message.reply_text("لا يوجد عملاء مسجلون.")
        return

    await send_long(
        update,
        "قائمة العملاء:\n",
        [f"{cid}. {name} — الدين: {money(debt)}" for cid, name, debt in rows],
    )


async def debts_list(update: Update, context: ContextTypes.DEFAULT_TYPE):
    debtors = [(name, debt) for _, name, debt in all_customer_debts() if debt > 0]

    if not debtors:
        await update.message.reply_text("لا توجد ديون حالية.")
        return

    total = sum(debt for _, debt in debtors)
    lines = [f"{name}: {money(debt)}" for name, debt in debtors]
    lines.append(f"\nإجمالي الديون: {money(total)}")

    await send_long(update, "العملاء المدينون:\n", lines)


# ============================================================
# بحث عن عميل
# ============================================================

async def search_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("اكتب اسم العميل:", reply_markup=CANCEL_KEYBOARD)
    return SEARCH_NAME


async def search_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    matches = find_customers(update.message.text)

    if not matches:
        await update.message.reply_text("العميل غير موجود.", reply_markup=MAIN_KEYBOARD)
        return ConversationHandler.END

    if len(matches) > 1:
        await update.message.reply_text(
            "وجدت أكثر من عميل، اكتب الاسم كاملًا:\n\n"
            + "\n".join(r["name"] for r in matches)
        )
        return SEARCH_NAME

    customer = matches[0]

    connection = get_connection()
    try:
        stats = connection.execute(
            """
            SELECT
                COUNT(*) AS invoices,
                COALESCE(SUM(total_amount), 0) AS total_sales,
                COALESCE(SUM(paid_amount), 0) AS total_paid
            FROM sales WHERE customer_id = ?
            """,
            (customer["id"],),
        ).fetchone()

        payments = connection.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM payments WHERE customer_id = ?",
            (customer["id"],),
        ).fetchone()[0]

        last = connection.execute(
            """
            SELECT id, created_at, total_amount, payment_method
            FROM sales WHERE customer_id = ?
            ORDER BY id DESC LIMIT 5
            """,
            (customer["id"],),
        ).fetchall()
    finally:
        connection.close()

    lines = [
        f"بيانات العميل: {customer['name']}",
        "",
        f"عدد الفواتير: {stats['invoices']}",
        f"إجمالي المشتريات: {money(stats['total_sales'])}",
        f"المدفوع وقت الشراء: {money(stats['total_paid'])}",
        f"تسديدات الديون: {money(payments)}",
        f"الدين الحالي: {money(get_customer_debt(customer['id']))}",
    ]

    if last:
        lines.append("\nآخر الفواتير:")
        for row in last:
            lines.append(
                f"رقم {row['id']} — {row['created_at'][:10]} — "
                f"{money(row['total_amount'])} — {row['payment_method']}"
            )

    await update.message.reply_text("\n".join(lines), reply_markup=MAIN_KEYBOARD)
    return ConversationHandler.END


# ============================================================
# التقارير
# ============================================================

def period_info(kind):
    current = datetime.now(TZ)
    if kind == "day":
        return "DATE({c}) = ?", current.strftime("%Y-%m-%d"), "تقرير اليوم"
    return "strftime('%Y-%m', {c}) = ?", current.strftime("%Y-%m"), "تقرير الشهر"


async def report(update, kind):
    condition, param, title = period_info(kind)
    cond_sales = condition.format(c="s.created_at")
    cond_plain = condition.format(c="created_at")

    connection = get_connection()
    try:
        debts = connection.execute(
            f"""
            SELECT COUNT(*) AS count, COALESCE(SUM(d.amount), 0) AS amount
            FROM debt_details d JOIN sales s ON s.id = d.sale_id
            WHERE {cond_sales}
            """,
            (param,),
        ).fetchone()
        entries = connection.execute(
            f"""
            SELECT COALESCE(c.name, 'عميل محذوف') AS customer,
                   d.product_description AS description, d.amount, d.created_at
            FROM debt_details d
            JOIN sales s ON s.id = d.sale_id
            LEFT JOIN customers c ON c.id = s.customer_id
            WHERE {cond_sales}
            ORDER BY d.created_at, d.sale_id
            """,
            (param,),
        ).fetchall()
        payments = connection.execute(
            f"SELECT COALESCE(SUM(amount), 0) FROM payments WHERE {cond_plain}",
            (param,),
        ).fetchone()[0]
    finally:
        connection.close()

    outstanding = sum(debt for _, _, debt in all_customer_debts())
    lines = [
        f"{title} — {param}",
        "",
        f"عدد الديون المسجلة: {debts['count']}",
        f"قيمة الديون المسجلة: {money(debts['amount'])}",
        f"التسديدات المسجلة: {money(payments)}",
        f"إجمالي الديون القائمة حاليًا: {money(outstanding)}",
    ]
    if entries:
        lines.append(chr(10) + "تفاصيل الديون:")
        for row in entries:
            lines.append(
                f"{row['customer']} — {row['description']} — {money(row['amount'])} — {row['created_at']}"
            )

    await send_long(update, "", lines)


async def today_report(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await report(update, "day")


async def month_report(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await report(update, "month")


# ============================================================
# النسخ الاحتياطي (يدوي + يومي تلقائي)
# ============================================================

def build_excel(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    connection = get_connection()
    try:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "الملخص"

        current_debt = sum(d for _, _, d in all_customer_debts())
        total_registered = connection.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM debt_details"
        ).fetchone()[0]
        total_payments = connection.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM payments"
        ).fetchone()[0]

        for row in [
            [f"دفتر ديون {SHOP_NAME}", ""],
            ["تاريخ إنشاء النسخة", now()],
            ["إجمالي الديون المسجلة تاريخيًا", total_registered],
            ["إجمالي التسديدات", total_payments],
            ["إجمالي الديون الحالية", current_debt],
        ]:
            summary.append(row)

        sheets = [
            (
                "العملاء والرصيد",
                ["رقم العميل", "اسم العميل", "الدين الحالي"],
                None,
            ),
            (
                "سجل الديون",
                ["رقم السجل", "التاريخ", "العميل", "المنتج أو الوصف", "المبلغ"],
                """
                SELECT s.id, d.created_at, COALESCE(c.name, 'عميل محذوف'),
                       d.product_description, d.amount
                FROM debt_details d
                JOIN sales s ON s.id = d.sale_id
                LEFT JOIN customers c ON c.id = s.customer_id
                ORDER BY d.created_at, s.id
                """,
            ),
            (
                "التسديدات",
                ["رقم التسديد", "التاريخ", "العميل", "المبلغ", "الطريقة"],
                """
                SELECT p.id, p.created_at, COALESCE(c.name, 'عميل محذوف'),
                       p.amount, p.payment_method
                FROM payments p LEFT JOIN customers c ON c.id = p.customer_id
                ORDER BY p.created_at, p.id
                """,
            ),
        ]

        for title, headers, sql in sheets:
            sheet = workbook.create_sheet(title)
            sheet.append(headers)
            if sql is None:
                for cid, name, debt in all_customer_debts():
                    sheet.append([cid, name, debt])
            else:
                for row in connection.execute(sql):
                    sheet.append(list(tuple(row)))

        for sheet in workbook.worksheets:
            sheet.sheet_view.rightToLeft = True
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
            for column_cells in sheet.columns:
                longest = max(
                    len("" if c.value is None else str(c.value)) for c in column_cells
                )
                sheet.column_dimensions[
                    get_column_letter(column_cells[0].column)
                ].width = min(max(longest + 2, 12), 40)

        workbook.save(path)
    finally:
        connection.close()


def make_backup_files():
    stamp = datetime.now(TZ).strftime("%Y%m%d_%H%M%S")
    folder = tempfile.mkdtemp(prefix="shop_backup_")
    db_copy = os.path.join(folder, f"shop_backup_{stamp}.db")
    excel = os.path.join(folder, f"shop_report_{stamp}.xlsx")

    source = get_connection()
    target = sqlite3.connect(db_copy)
    try:
        with target:
            source.backup(target)
    finally:
        target.close()
        source.close()

    check = sqlite3.connect(db_copy)
    try:
        integrity = check.execute("PRAGMA integrity_check").fetchone()[0]
        tables = {
            r[0]
            for r in check.execute("SELECT name FROM sqlite_master WHERE type = 'table'")
        }
    finally:
        check.close()

    if integrity != "ok":
        raise RuntimeError(f"فشل فحص سلامة القاعدة: {integrity}")

    required = {"customers", "sales", "payments", "debt_details"}
    if not required.issubset(tables):
        raise RuntimeError("النسخة لا تحتوي على الجداول الأساسية.")

    build_excel(excel)

    return folder, db_copy, excel


async def send_backup(bot, chat_id, label="النسخة الاحتياطية"):
    folder = None
    try:
        folder, db_copy, excel = make_backup_files()

        with open(db_copy, "rb") as file:
            await bot.send_document(
                chat_id=chat_id,
                document=file,
                filename=os.path.basename(db_copy),
                caption=f"{label}: قاعدة البيانات.\nفحص السلامة: سليم",
            )

        with open(excel, "rb") as file:
            await bot.send_document(
                chat_id=chat_id,
                document=file,
                filename=os.path.basename(excel),
                caption=f"{label}: تقرير Excel.",
            )

        return True

    except Exception as error:
        logger.exception("backup failed")
        await bot.send_message(chat_id=chat_id, text=f"فشل النسخ الاحتياطي: {error}")
        return False

    finally:
        if folder:
            shutil.rmtree(folder, ignore_errors=True)


async def backup_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await send_backup(context.bot, update.effective_chat.id)


async def scheduled_backup(context: ContextTypes.DEFAULT_TYPE):
    await send_backup(context.bot, context.job.data, label="نسخة يومية تلقائية")


# ============================================================
# حذف عميل (لا يُحذف من عليه دين، والفواتير تبقى في التقارير)
# ============================================================

async def delete_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text(
        "اكتب اسم العميل المراد حذفه.\nلا يمكن حذف عميل عليه دين.",
        reply_markup=CANCEL_KEYBOARD,
    )
    return DELETE_NAME


async def delete_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    matches = find_customers(update.message.text)

    if not matches:
        await update.message.reply_text("العميل غير موجود.", reply_markup=MAIN_KEYBOARD)
        return ConversationHandler.END

    if len(matches) > 1:
        await update.message.reply_text(
            "وجدت أكثر من عميل، اكتب الاسم كاملًا:\n\n"
            + "\n".join(r["name"] for r in matches)
        )
        return DELETE_NAME

    customer = matches[0]
    debt = get_customer_debt(customer["id"])

    if debt > 0:
        await update.message.reply_text(
            f"لا يمكن حذف {customer['name']}.\n"
            f"دينه الحالي: {money(debt)}\n"
            "يجب تسديد الدين بالكامل أولًا.",
            reply_markup=MAIN_KEYBOARD,
        )
        return ConversationHandler.END

    context.user_data["delete_id"] = customer["id"]
    context.user_data["delete_name"] = customer["name"]

    await update.message.reply_text(
        f"العميل: {customer['name']}\n"
        "سيُحذف من قائمة العملاء، وتبقى سجلات ديونه وتسديداته دون اسم.\n"
        "سأرسل لك نسخة احتياطية قبل الحذف.\n\n"
        "للتأكيد اكتب: نعم\nللإلغاء اكتب: لا"
    )
    return DELETE_CONFIRM


async def delete_confirm(update: Update, context: ContextTypes.DEFAULT_TYPE):
    answer = update.message.text.strip()

    if answer == "لا":
        context.user_data.clear()
        await update.message.reply_text("تم إلغاء الحذف.", reply_markup=MAIN_KEYBOARD)
        return ConversationHandler.END

    if answer != "نعم":
        await update.message.reply_text("اكتب نعم للتأكيد أو لا للإلغاء.")
        return DELETE_CONFIRM

    customer_id = context.user_data.get("delete_id")
    customer_name = context.user_data.get("delete_name")

    backed_up = await send_backup(
        context.bot, update.effective_chat.id, label="نسخة أمان قبل الحذف"
    )

    if not backed_up:
        await update.message.reply_text(
            "أوقفت الحذف لأن نسخة الأمان فشلت.", reply_markup=MAIN_KEYBOARD
        )
        context.user_data.clear()
        return ConversationHandler.END

    connection = get_connection()
    try:
        cursor = connection.cursor()
        cursor.execute("BEGIN IMMEDIATE")

        owed = cursor.execute(
            """
            SELECT COALESCE((SELECT SUM(debt_amount) FROM sales WHERE customer_id = ?), 0)
                 - COALESCE((SELECT SUM(amount) FROM payments WHERE customer_id = ?), 0)
            """,
            (customer_id, customer_id),
        ).fetchone()[0]

        if owed > 0:
            connection.rollback()
            await update.message.reply_text(
                "تم إيقاف الحذف لأن العميل عليه دين الآن.",
                reply_markup=MAIN_KEYBOARD,
            )
        else:
            cursor.execute(
                "UPDATE sales SET customer_id = NULL WHERE customer_id = ?",
                (customer_id,),
            )
            cursor.execute(
                "UPDATE payments SET customer_id = NULL WHERE customer_id = ?",
                (customer_id,),
            )
            cursor.execute("DELETE FROM customers WHERE id = ?", (customer_id,))
            connection.commit()
            await update.message.reply_text(
                f"تم حذف العميل {customer_name}.", reply_markup=MAIN_KEYBOARD
            )
    except Exception as error:
        connection.rollback()
        logger.exception("delete failed")
        await update.message.reply_text(
            f"فشل الحذف: {error}", reply_markup=MAIN_KEYBOARD
        )
    finally:
        connection.close()
        context.user_data.clear()

    return ConversationHandler.END


# ============================================================
# الأوامر العامة
# ============================================================

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        f"أهلاً بك في {SHOP_NAME}.\n\nاختر العملية:",
        reply_markup=MAIN_KEYBOARD,
    )


async def myid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if user:
        await update.message.reply_text(f"Telegram User ID:\n\n{user.id}")


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text(
        "تم إلغاء العملية الحالية. اضغط الزر المطلوب للبدء من جديد.",
        reply_markup=MAIN_KEYBOARD,
    )
    return ConversationHandler.END


async def unknown_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("اختر من القائمة:", reply_markup=MAIN_KEYBOARD)


async def deny(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.message:
        await update.message.reply_text("هذا البوت مخصص لأصحاب المحل فقط.")


async def on_error(update, context: ContextTypes.DEFAULT_TYPE):
    logger.error("Unhandled error", exc_info=context.error)


# ============================================================
# التشغيل
# ============================================================

def button(text):
    return filters.Regex("^" + re.escape(text) + "$") & ALLOWED


def text_state(handler):
    return [MessageHandler(TEXT_STATE, handler)]


def run_render_webhook(application):
    """تشغيل البوت عبر Webhook على خدمة Render Web Service."""
    base_url = (os.getenv("WEBHOOK_URL") or os.getenv("RENDER_EXTERNAL_URL") or "").strip().rstrip("/")
    if not base_url.startswith("https://"):
        raise RuntimeError("WEBHOOK_URL أو RENDER_EXTERNAL_URL يجب أن يبدأ بـ https://")
    secret = os.getenv("TELEGRAM_WEBHOOK_SECRET", "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,256}", secret):
        raise RuntimeError("اضبط TELEGRAM_WEBHOOK_SECRET على رمز آمن صالح.")
    path = (os.getenv("WEBHOOK_PATH", "telegram").strip().strip("/") or "telegram")
    port = int(os.getenv("PORT", "10000"))
    application.run_webhook(
        listen="0.0.0.0", port=port, url_path=path,
        webhook_url=f"{base_url}/{path}", secret_token=secret,
        bootstrap_retries=-1,
    )


def conversation(entry_text, entry_handler, states):
    return ConversationHandler(
        entry_points=[MessageHandler(button(entry_text), entry_handler)],
        states=states,
        fallbacks=[
            CommandHandler("cancel", cancel),
            MessageHandler(MENU_FILTER, cancel),
        ],
    )


def main():
    if not TOKEN:
        raise RuntimeError("TELEGRAM_TOKEN غير موجود في متغيرات البيئة.")

    init_database()
    ensure_debt_details_table()

    application = Application.builder().token(TOKEN).build()
    application.add_error_handler(on_error)
    application.add_handler(CommandHandler("myid", myid))

    if not ALLOWED_IDS:
        logger.warning(
            "OWNER_ID غير مضبوط. البوت في وضع الإعداد: أرسل /myid ثم ضع الرقم في OWNER_ID."
        )
        run_render_webhook(application)
        return

    application.add_handler(CommandHandler("start", start, filters=ALLOWED))

    application.add_handler(conversation(BTN_DEBT_NEW, new_debt_start, {
        NEW_DEBT_NAME: text_state(new_debt_name),
        NEW_DEBT_ITEM: text_state(new_debt_item),
        NEW_DEBT_AMOUNT: text_state(new_debt_amount),
    }))

    application.add_handler(conversation(BTN_DEBT_PAY, debt_start, {
        DEBT_NAME: text_state(debt_name),
        DEBT_AMOUNT: text_state(debt_amount),
        DEBT_METHOD: text_state(debt_method),
    }))

    application.add_handler(conversation(BTN_DELETE, delete_start, {
        DELETE_NAME: text_state(delete_name),
        DELETE_CONFIRM: text_state(delete_confirm),
    }))

    application.add_handler(MessageHandler(button(BTN_CUSTOMERS), customers_list))
    application.add_handler(MessageHandler(button(BTN_DEBTS), debts_list))
    application.add_handler(MessageHandler(button(BTN_TODAY), today_report))
    application.add_handler(MessageHandler(button(BTN_MONTH), month_report))
    application.add_handler(MessageHandler(button(BTN_BACKUP), backup_command))
    application.add_handler(MessageHandler(button(BTN_CANCEL), cancel))

    application.add_handler(
        MessageHandler(ALLOWED & filters.TEXT & ~filters.COMMAND, unknown_text)
    )
    application.add_handler(MessageHandler(~ALLOWED, deny))

    if application.job_queue:
        application.job_queue.run_daily(
            scheduled_backup,
            time=time(hour=BACKUP_HOUR, minute=0, tzinfo=TZ),
            data=ALLOWED_IDS[0],
            name="daily_backup",
        )
    else:
        logger.warning("JobQueue غير متاح: ثبّت python-telegram-bot[job-queue].")

    logger.info("%s bot is running with Telegram webhook...", SHOP_NAME)
    run_render_webhook(application)


if __name__ == "__main__":
    main()
